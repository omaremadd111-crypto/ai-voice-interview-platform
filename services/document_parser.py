"""Document text extraction: PDF, DOCX, TXT, MD -> normalized ParsedDocument.

Extracts once per document and never leaks a library-specific exception
(pypdf/python-docx) to callers -- everything funnels through DocumentParseError
or UnsupportedDocumentError.
"""
import re
import unicodedata
from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from pydantic import BaseModel
from pypdf import PdfReader

from config.settings import Settings


class UnsupportedDocumentError(Exception):
    """Raised when a document's file extension is not supported."""


class DocumentParseError(Exception):
    """Raised when a document cannot be parsed into usable text."""


class ParsedDocument(BaseModel):
    filename: str
    extension: str
    text: str
    char_count: int
    page_count: int | None = None
    truncated: bool = False


_CONTROL_CHAR_PATTERN = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]")
_INLINE_WHITESPACE_PATTERN = re.compile(r"[ \t]+")
_BLANK_LINE_RUN_PATTERN = re.compile(r"\n{3,}")


def _normalize_text(raw: str) -> str:
    text = unicodedata.normalize("NFKC", raw)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _CONTROL_CHAR_PATTERN.sub("", text)
    text = _INLINE_WHITESPACE_PATTERN.sub(" ", text)
    text = _BLANK_LINE_RUN_PATTERN.sub("\n\n", text)
    lines = [line.strip() for line in text.split("\n")]
    return "\n".join(lines).strip()


class DocumentParser:
    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}

    def __init__(self, settings: Settings | None = None) -> None:
        self._max_chars = (settings or Settings()).max_doc_chars

    def parse_path(self, path: str | Path) -> ParsedDocument:
        path = Path(path)
        try:
            data = path.read_bytes()
        except OSError as exc:
            raise DocumentParseError(
                f"Could not read document '{path.name}': file is missing or unreadable"
            ) from exc
        return self.parse_bytes(data, filename=path.name)

    def parse_bytes(self, data: bytes, filename: str) -> ParsedDocument:
        extension = Path(filename).suffix.lower()
        if extension not in self.SUPPORTED_EXTENSIONS:
            raise UnsupportedDocumentError(
                f"Unsupported file type '{extension}'. Supported types: "
                f"{', '.join(sorted(self.SUPPORTED_EXTENSIONS))}"
            )

        if extension == ".pdf":
            raw_text, page_count = self._extract_pdf(data, filename)
        elif extension == ".docx":
            raw_text, page_count = self._extract_docx(data, filename)
        else:
            raw_text, page_count = self._extract_plain_text(data, filename)

        text = _normalize_text(raw_text)
        if not text:
            raise DocumentParseError("no extractable text")

        truncated = len(text) > self._max_chars
        if truncated:
            text = text[: self._max_chars]

        return ParsedDocument(
            filename=filename,
            extension=extension,
            text=text,
            char_count=len(text),
            page_count=page_count,
            truncated=truncated,
        )

    def _extract_plain_text(self, data: bytes, filename: str) -> tuple[str, None]:
        try:
            return data.decode("utf-8"), None
        except UnicodeDecodeError:
            try:
                return data.decode("latin-1"), None
            except UnicodeDecodeError as exc:
                raise DocumentParseError(f"Could not decode '{filename}' as text") from exc

    def _extract_pdf(self, data: bytes, filename: str) -> tuple[str, int]:
        try:
            reader = PdfReader(BytesIO(data))
            page_count = len(reader.pages)
            pages_text = [page.extract_text() or "" for page in reader.pages]
        except Exception as exc:
            raise DocumentParseError(f"Could not read PDF '{filename}': corrupted or invalid file") from exc
        return "\n".join(pages_text), page_count

    def _extract_docx(self, data: bytes, filename: str) -> tuple[str, None]:
        try:
            doc = DocxDocument(BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        parts.append(cell.text)
        except Exception as exc:
            raise DocumentParseError(f"Could not read DOCX '{filename}': corrupted or invalid file") from exc
        return "\n".join(parts), None
