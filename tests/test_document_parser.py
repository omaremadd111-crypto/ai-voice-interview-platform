"""Document parser tests: extraction, normalization, truncation, error handling."""
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from config.settings import Settings
from services.document_parser import (
    DocumentParseError,
    DocumentParser,
    ParsedDocument,
    UnsupportedDocumentError,
)


def _docx_bytes(paragraphs: list[str]) -> bytes:
    doc = DocxDocument()
    for para in paragraphs:
        doc.add_paragraph(para)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def parser() -> DocumentParser:
    return DocumentParser(settings=Settings(max_doc_chars=20000))


# ---- TXT ----

def test_txt_extraction(parser: DocumentParser) -> None:
    data = "Line one.\nLine two.\nRelevant skills: Python, SQL.".encode("utf-8")
    result = parser.parse_bytes(data, filename="notes.txt")
    assert isinstance(result, ParsedDocument)
    assert result.filename == "notes.txt"
    assert result.extension == ".txt"
    assert "Line one." in result.text
    assert "Relevant skills: Python, SQL." in result.text
    assert result.char_count == len(result.text)
    assert result.page_count is None
    assert result.truncated is False


def test_txt_extraction_latin1_fallback(parser: DocumentParser) -> None:
    # Not valid UTF-8, but decodable as latin-1 -- must not raise.
    data = "Café résumé".encode("latin-1")
    result = parser.parse_bytes(data, filename="notes.txt")
    assert "Caf" in result.text


# ---- Markdown ----

def test_markdown_extraction(parser: DocumentParser) -> None:
    data = "# Job Title\n\n**Company:** FlairsTech\n\n- Python\n- SQL\n".encode("utf-8")
    result = parser.parse_bytes(data, filename="job.md")
    assert result.extension == ".md"
    assert "Job Title" in result.text
    assert "FlairsTech" in result.text
    assert result.page_count is None


# ---- DOCX ----

def test_docx_extraction(parser: DocumentParser) -> None:
    data = _docx_bytes(["Senior AI Engineer", "5 years of experience with RAG systems."])
    result = parser.parse_bytes(data, filename="cv.docx")
    assert result.extension == ".docx"
    assert "Senior AI Engineer" in result.text
    assert "RAG systems" in result.text
    assert result.page_count is None
    assert result.char_count == len(result.text)


def test_docx_extraction_includes_table_text(parser: DocumentParser) -> None:
    doc = DocxDocument()
    doc.add_paragraph("Skills Matrix")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "Expert"
    buf = BytesIO()
    doc.save(buf)
    result = parser.parse_bytes(buf.getvalue(), filename="cv_with_table.docx")
    assert "Python" in result.text
    assert "Expert" in result.text


# ---- PDF ----

def test_pdf_extraction_from_bytes(parser: DocumentParser, fixtures_dir: Path) -> None:
    data = (fixtures_dir / "sample.pdf").read_bytes()
    result = parser.parse_bytes(data, filename="sample.pdf")
    assert result.extension == ".pdf"
    assert result.page_count == 1
    assert "DEMO DATA" in result.text
    assert "PDF text extraction" in result.text
    assert result.char_count == len(result.text)


def test_pdf_extraction_from_path(parser: DocumentParser, fixtures_dir: Path) -> None:
    result = parser.parse_path(fixtures_dir / "sample.pdf")
    assert result.filename == "sample.pdf"
    assert result.page_count == 1
    assert "DEMO DATA" in result.text


def test_missing_path_raises_document_parse_error(parser: DocumentParser, tmp_path: Path) -> None:
    with pytest.raises(DocumentParseError) as exc_info:
        parser.parse_path(tmp_path / "missing.txt")
    assert type(exc_info.value) is DocumentParseError
    assert "missing or unreadable" in str(exc_info.value)


def test_unreadable_path_raises_document_parse_error(
    parser: DocumentParser, monkeypatch: pytest.MonkeyPatch,
) -> None:
    def _raise_permission_error(_path: Path) -> bytes:
        raise PermissionError("access denied")

    monkeypatch.setattr(Path, "read_bytes", _raise_permission_error)
    with pytest.raises(DocumentParseError) as exc_info:
        parser.parse_path(Path("private.txt"))
    assert type(exc_info.value) is DocumentParseError
    assert "missing or unreadable" in str(exc_info.value)


# ---- unsupported extension ----

def test_unsupported_extension_raises(parser: DocumentParser) -> None:
    with pytest.raises(UnsupportedDocumentError):
        parser.parse_bytes(b"whatever", filename="resume.exe")


def test_unsupported_extension_message_lists_supported_types(parser: DocumentParser) -> None:
    with pytest.raises(UnsupportedDocumentError) as exc_info:
        parser.parse_bytes(b"whatever", filename="resume.png")
    message = str(exc_info.value)
    assert ".pdf" in message and ".docx" in message and ".txt" in message


# ---- corrupted files: must raise DocumentParseError, never a raw library exception ----

def test_corrupted_pdf_raises_document_parse_error(parser: DocumentParser) -> None:
    with pytest.raises(DocumentParseError) as exc_info:
        parser.parse_bytes(b"this is not a valid PDF file at all", filename="broken.pdf")
    assert type(exc_info.value) is DocumentParseError
    assert "PdfRead" not in str(exc_info.value)


def test_corrupted_docx_raises_document_parse_error(parser: DocumentParser) -> None:
    with pytest.raises(DocumentParseError) as exc_info:
        parser.parse_bytes(b"this is not a valid DOCX file at all", filename="broken.docx")
    assert type(exc_info.value) is DocumentParseError


# ---- empty extracted text ----

def test_empty_text_raises_exact_message(parser: DocumentParser) -> None:
    with pytest.raises(DocumentParseError) as exc_info:
        parser.parse_bytes(b"   \n\n\t  \x00\x01  ", filename="blank.txt")
    assert str(exc_info.value) == "no extractable text"


# ---- normalization ----

def test_normalization_strips_control_chars_and_collapses_whitespace(parser: DocumentParser) -> None:
    data = "Title\x00\x01\n\n\n\nToo   many    spaces.\n\n\n\nEnd.".encode("utf-8")
    result = parser.parse_bytes(data, filename="messy.txt")
    assert "\x00" not in result.text
    assert "\x01" not in result.text
    assert "   " not in result.text  # no runs of 3+ spaces survive
    assert "\n\n\n" not in result.text  # no runs of 3+ newlines survive
    assert "Title" in result.text and "End." in result.text


# ---- truncation ----

def test_truncation_sets_flag_and_enforces_max_chars() -> None:
    small_parser = DocumentParser(settings=Settings(max_doc_chars=50))
    data = ("A" * 5000).encode("utf-8")
    result = small_parser.parse_bytes(data, filename="long.txt")
    assert result.truncated is True
    assert result.char_count == 50
    assert result.text == "A" * 50


def test_no_truncation_when_under_limit() -> None:
    small_parser = DocumentParser(settings=Settings(max_doc_chars=5000))
    data = "short text".encode("utf-8")
    result = small_parser.parse_bytes(data, filename="short.txt")
    assert result.truncated is False
    assert result.char_count == len("short text")


# ---- default settings wiring ----

def test_default_constructor_uses_settings_defaults() -> None:
    default_parser = DocumentParser()
    data = "hello".encode("utf-8")
    result = default_parser.parse_bytes(data, filename="hello.txt")
    assert result.truncated is False
